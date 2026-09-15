from flask import Flask, render_template, request, jsonify, send_from_directory, redirect, url_for, session, flash, Response, abort
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import simpleSplit
from pypdf import PdfReader, PdfWriter
import io as _io
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import inspect, text
from datetime import datetime
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
import os, shutil, json, functools, base64, unicodedata
import requests as http_requests
from urllib.parse import quote, unquote, urlparse
import re
import openpyxl
import xlrd

app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'genmar-fatura-2026')
app.config['JSON_AS_ASCII'] = False
app.json.ensure_ascii = False
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

@app.after_request
def turkce_karakter_kodlamasi(response):
    """HTML, metin ve JSON cevaplarini her tarayicida UTF-8 olarak goster."""
    if response.mimetype in ('text/html', 'text/plain', 'application/json'):
        response.content_type = f'{response.mimetype}; charset=utf-8'
    return response

# Veritabanı: Bulutta PostgreSQL, lokalde SQLite
DATABASE_URL = os.environ.get('DATABASE_URL', '')
if DATABASE_URL and 'pg8000' not in DATABASE_URL and '+' not in DATABASE_URL.split('://')[0]:
    DATABASE_URL = DATABASE_URL.replace('postgresql://', 'postgresql+pg8000://', 1).replace('postgres://', 'postgresql+pg8000://', 1)
if DATABASE_URL.startswith('postgresql://') and 'psycopg' not in DATABASE_URL:
    DATABASE_URL = DATABASE_URL.replace('postgresql://', 'postgresql+psycopg2://', 1)
app.config['SQLALCHEMY_DATABASE_URI'] = DATABASE_URL or f'sqlite:///{os.path.join(BASE_DIR, "database", "fatura.db")}'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['UPLOAD_FOLDER'] = os.path.join(BASE_DIR, 'static', 'uploads')
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024

# Supabase Storage konfigürasyonu
SUPABASE_URL = os.environ.get('SUPABASE_URL', '')
SUPABASE_KEY = os.environ.get('SUPABASE_KEY', '')
SUPABASE_BUCKET = os.environ.get('SUPABASE_BUCKET', 'faturalar')
BULUT_MOD = bool(SUPABASE_URL and SUPABASE_KEY)

# Eski faturalar Cloudinary'de tutuluyor.
CLOUDINARY_CLOUD_NAME = os.environ.get('CLOUDINARY_CLOUD_NAME', '')
CLOUDINARY_API_KEY = os.environ.get('CLOUDINARY_API_KEY', '')
CLOUDINARY_API_SECRET = os.environ.get('CLOUDINARY_API_SECRET', '')

def supabase_yukle(dosya_bytes, dosya_yolu):
    """Supabase Storage'a dosya yükle, public URL döndür."""
    url = f"{SUPABASE_URL}/storage/v1/object/{SUPABASE_BUCKET}/{dosya_yolu}"
    headers = {
        'Authorization': f'Bearer {SUPABASE_KEY}',
        'Content-Type': 'application/pdf',
        'x-upsert': 'true'
    }
    r = http_requests.post(url, headers=headers, data=dosya_bytes)
    if r.status_code in (200, 201):
        return f"{SUPABASE_URL}/storage/v1/object/public/{SUPABASE_BUCKET}/{dosya_yolu}"
    raise Exception(f"Supabase yukleme hatasi: {r.status_code} {r.text}")

def supabase_indir(public_url):
    """Supabase dosyasını public veya private bucket'tan indir."""
    dosya_yolu = supabase_dosya_yolu(public_url)
    encoded_path = quote(unquote(dosya_yolu), safe='/')
    private_url = f"{SUPABASE_URL}/storage/v1/object/{SUPABASE_BUCKET}/{encoded_path}"
    headers = {
        'Authorization': f'Bearer {SUPABASE_KEY}',
        'apikey': SUPABASE_KEY,
    }
    r = http_requests.get(private_url, headers=headers, timeout=30)
    if r.status_code == 200:
        return r.content

    # Eski public kayıtlarla geriye dönük uyumluluk.
    if public_url.startswith(('http://', 'https://')):
        r = http_requests.get(public_url, timeout=30)
        if r.status_code == 200:
            return r.content
    raise Exception(f"Supabase indirme hatasi: {r.status_code}")

def cloudinary_indir(public_url):
    """Eski Cloudinary PDF'lerini kimlik doğrulamalı, süreli URL ile indir."""
    if not (CLOUDINARY_CLOUD_NAME and CLOUDINARY_API_KEY and CLOUDINARY_API_SECRET):
        raise Exception('Cloudinary kimlik bilgileri eksik')

    import cloudinary
    import cloudinary.utils
    cloudinary.config(
        cloud_name=CLOUDINARY_CLOUD_NAME,
        api_key=CLOUDINARY_API_KEY,
        api_secret=CLOUDINARY_API_SECRET,
        secure=True,
    )

    yol = unquote(urlparse(public_url).path)
    eslesme = re.search(
        r'/raw/(upload|private|authenticated)/(?:s--[^/]+--/)?(?:v(\d+)/)?(.+)$', yol
    )
    if not eslesme:
        raise Exception('Geçersiz Cloudinary PDF adresi')

    delivery_type, version, public_id_tam = eslesme.groups()
    public_id_uzantisiz = (
        public_id_tam[:-4] if public_id_tam.lower().endswith('.pdf')
        else public_id_tam
    )

    # Cloudinary "raw" dosyalarında public_id çoğunlukla uzantıyı da içerir.
    # Eski kayıtların bir kısmı ise uzantısız kimlikle yüklenmiş olabileceğinden
    # iki biçimi de güvenli, imzalı indirme bağlantısıyla deniyoruz.
    # Kısıtlı PDF teslimatında imzalı delivery URL en güvenilir yöntemdir.
    delivery_url, _ = cloudinary.utils.cloudinary_url(
        public_id_tam,
        resource_type='raw',
        type=delivery_type,
        version=version,
        sign_url=True,
        secure=True,
    )
    delivery_cevap = http_requests.get(delivery_url, timeout=30)
    if delivery_cevap.status_code == 200 and delivery_cevap.content:
        return delivery_cevap.content

    denemeler = [
        (public_id_tam, None),
        (public_id_uzantisiz, 'pdf'),
    ]
    durumlar = [str(delivery_cevap.status_code)]
    for public_id, dosya_formati in denemeler:
        signed_url = cloudinary.utils.private_download_url(
            public_id,
            dosya_formati,
            resource_type='raw',
            type=delivery_type,
            attachment=False,
        )
        r = http_requests.get(signed_url, timeout=30)
        if r.status_code == 200 and r.content:
            return r.content
        durumlar.append(str(r.status_code))

    raise Exception(f"Cloudinary indirme hatasi: {', '.join(durumlar)}")

db = SQLAlchemy(app)

# ─── LOGIN DECORATOR ────────────────────────────────────────────────────────

def login_gerekli(f):
    @functools.wraps(f)
    def decorated(*args, **kwargs):
        if 'kullanici_id' not in session:
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated

# ─── PDF DAMGA ──────────────────────────────────────────────────────────────

def pdf_damga_ekle(pdf_path, output_path, bilgiler):
    try:
        reader = PdfReader(pdf_path)
        writer = PdfWriter()
        for page in reader.pages:
            w = float(page.mediabox.width)
            h = float(page.mediabox.height)
            overlay_buf = _io.BytesIO()
            c = rl_canvas.Canvas(overlay_buf, pagesize=(w, h))
            bx, by, bw, bh = w - 205, h - 135, 190, 120
            c.setFillColorRGB(1.0, 0.94, 0.94)
            c.roundRect(bx, by, bw, bh, 5, fill=1, stroke=0)
            c.setLineWidth(1.5); c.setStrokeColorRGB(0.82, 0.08, 0.08)
            c.roundRect(bx, by, bw, bh, 5, fill=0, stroke=1)
            c.setFillColorRGB(0.82, 0.08, 0.08)
            c.roundRect(bx, by+bh-18, bw, 18, 5, fill=1, stroke=0)
            c.rect(bx, by+bh-18, bw, 9, fill=1, stroke=0)
            c.setFillColorRGB(1,1,1); c.setFont("Helvetica-Bold", 8)
            c.drawCentredString(bx+bw/2, by+bh-13, "ONAYLANDI")
            def satir(label, deger, yo):
                c.setFillColorRGB(0.5,0.05,0.05); c.setFont("Helvetica-Bold",7)
                c.drawString(bx+8, by+bh-yo, label)
                c.setFillColorRGB(0.15,0.15,0.15); c.setFont("Helvetica",7)
                d = str(deger) if deger else '-'
                if len(d) > 30: d = d[:28]+'..'
                c.drawString(bx+72, by+bh-yo, d)
            projeler = bilgiler.get('projeler', [])
            if projeler:
                satir("PROJELER:", ', '.join([p.get('ana','') for p in projeler]), 32)
                satir("ALT:", ', '.join([p.get('alt','') for p in projeler]), 44)
            else:
                satir("ANA PROJE:", bilgiler.get('ana_proje_kod','-'), 32)
                satir("ALT PROJE:", bilgiler.get('alt_proje_kod','-'), 44)
            satir("NOT:", bilgiler.get('not_alani','-'), 56)
            satir("ÖDEME:", bilgiler.get('odeme_notu','-'), 68)
            c.setStrokeColorRGB(0.82,0.08,0.08); c.setLineWidth(0.4)
            c.line(bx+6, by+bh-74, bx+bw-6, by+bh-74)
            satir("ONAYLAYAN:", bilgiler.get('onaylayan_ad','-'), 86)
            satir("DEPARTMAN:", bilgiler.get('departman_kod','-'), 98)
            satir("TARİH:", bilgiler.get('tarih','-'), 110)
            c.setFillColorRGB(0.82,0.08,0.08); c.setFont("Helvetica-Bold",11)
            c.drawRightString(bx+bw-8, by+6, bilgiler.get('kisaltma',''))
            c.save(); overlay_buf.seek(0)
            page.merge_page(PdfReader(overlay_buf).pages[0])
            writer.add_page(page)
        with open(output_path, 'wb') as f:
            writer.write(f)
        return True
    except Exception as e:
        print(f"PDF damga hatasi: {e}"); return False

def _damga_fontlarini_hazirla():
    if 'DamgaNormal' in pdfmetrics.getRegisteredFontNames():
        return 'DamgaNormal', 'DamgaKalin'
    import reportlab
    font_dir = os.path.join(os.path.dirname(reportlab.__file__), 'fonts')
    pdfmetrics.registerFont(TTFont('DamgaNormal', os.path.join(font_dir, 'Vera.ttf')))
    pdfmetrics.registerFont(TTFont('DamgaKalin', os.path.join(font_dir, 'VeraBd.ttf')))
    return 'DamgaNormal', 'DamgaKalin'

def pdf_damga_bytes(pdf_bytes, bilgiler):
    """Onay damgasini PDF'nin ilk sayfasina ekleyip yeni PDF baytlarini dondurur."""
    normal_font, kalin_font = _damga_fontlarini_hazirla()
    reader = PdfReader(_io.BytesIO(pdf_bytes))
    writer = PdfWriter()

    for sayfa_no, page in enumerate(reader.pages):
        if sayfa_no == 0:
            # Döndürülmüş PDF'lerde damganın görünür alanın dışında kalmasını önle.
            if int(page.get('/Rotate', 0) or 0) % 360:
                page.transfer_rotation_to_content()
            w = float(page.mediabox.width)
            h = float(page.mediabox.height)
            overlay_buf = _io.BytesIO()
            c = rl_canvas.Canvas(overlay_buf, pagesize=(w, h))
            bw = min(285, w * 0.48)
            bh = 178
            bx, by = w - bw - 18, h - bh - 18
            kirmizi = (0.88, 0.08, 0.04)

            c.setFillColorRGB(1.0, 0.96, 0.95)
            c.setStrokeColorRGB(*kirmizi)
            c.setLineWidth(2)
            c.roundRect(bx, by, bw, bh, 8, fill=1, stroke=1)
            c.setFillColorRGB(*kirmizi)
            c.roundRect(bx, by + bh - 28, bw, 28, 8, fill=1, stroke=0)
            c.rect(bx, by + bh - 28, bw, 12, fill=1, stroke=0)
            c.setFillColorRGB(1, 1, 1)
            c.setFont(kalin_font, 12)
            c.drawCentredString(bx + bw / 2, by + bh - 19, 'ONAYLANDI')

            label_x, deger_x = bx + 12, bx + 103
            satir_y = by + bh - 47
            satir_araligi = 16

            def satir(label, deger, y):
                c.setFillColorRGB(0.52, 0.04, 0.03)
                c.setFont(kalin_font, 8)
                c.drawString(label_x, y, label)
                c.setFillColorRGB(0.16, 0.16, 0.18)
                c.setFont(normal_font, 8)
                metin = metin_degeri(deger) or '-'
                parcalar = simpleSplit(metin, normal_font, 8, bw - 115)
                c.drawString(deger_x, y, parcalar[0] if parcalar else '-')

            projeler = bilgiler.get('projeler') or []
            ana = ', '.join(p.get('ana') or '-' for p in projeler) if projeler else '-'
            alt = ', '.join(p.get('alt') or '-' for p in projeler) if projeler else '-'
            satir('ANA PROJE:', ana, satir_y)
            satir('ALT PROJE:', alt, satir_y - satir_araligi)
            satir('ONAY NOTU:', bilgiler.get('not_alani'), satir_y - 2 * satir_araligi)
            satir('ÖDEME NOTU:', bilgiler.get('odeme_notu'), satir_y - 3 * satir_araligi)
            c.setStrokeColorRGB(*kirmizi)
            c.setLineWidth(0.6)
            c.line(bx + 10, satir_y - 3 * satir_araligi - 8,
                   bx + bw - 10, satir_y - 3 * satir_araligi - 8)
            satir('ONAYLAYAN:', bilgiler.get('onaylayan_ad'), satir_y - 4 * satir_araligi - 5)
            satir('DEPARTMAN:', bilgiler.get('departman_kod'), satir_y - 5 * satir_araligi - 5)
            satir('TARİH:', bilgiler.get('tarih'), satir_y - 6 * satir_araligi - 5)
            c.setFillColorRGB(*kirmizi)
            c.setFont(kalin_font, 13)
            c.drawRightString(bx + bw - 10, by + 9, bilgiler.get('kisaltma') or '')
            c.save()
            overlay_buf.seek(0)
            page.merge_page(PdfReader(overlay_buf).pages[0])
        writer.add_page(page)

    sonuc = _io.BytesIO()
    writer.add_metadata({'/GenmarOnayDamgasi': '1'})
    writer.write(sonuc)
    return sonuc.getvalue()

def pdf_damgali_mi(pdf_bytes):
    """PDF'nin daha önce onay damgası alıp almadığını güvenle belirler."""
    try:
        reader = PdfReader(_io.BytesIO(pdf_bytes))
        if (reader.metadata or {}).get('/GenmarOnayDamgasi') == '1':
            return True
        ilk_sayfa = reader.pages[0].extract_text() or '' if reader.pages else ''
        return 'ONAYLANDI' in ilk_sayfa.upper()
    except Exception:
        return False

# ─── MODELLER ───────────────────────────────────────────────────────────────

class Departman(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    kod = db.Column(db.String(10), unique=True, nullable=False)
    ad = db.Column(db.String(50), nullable=False)

class Kullanici(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    ad_soyad = db.Column(db.String(100), nullable=False)
    email = db.Column(db.String(254), unique=True, nullable=True)
    atanan_faturalari_gizle = db.Column(db.Boolean, default=False, nullable=False)
    sifre_hash = db.Column(db.String(200), default=None)
    departman_id = db.Column(db.Integer, db.ForeignKey('departman.id'))
    departman = db.relationship('Departman', backref='kullanicilar')
    @property
    def kisaltma(self):
        p = self.ad_soyad.strip().split()
        return (p[0][0]+p[-1][0]).upper() if len(p)>=2 else p[0][0].upper()
    def sifre_kontrol(self, sifre):
        if not self.sifre_hash:
            return sifre == '1'
        return check_password_hash(self.sifre_hash, sifre)
    def sifre_ayarla(self, sifre):
        self.sifre_hash = generate_password_hash(sifre)

class AnaProje(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    kod = db.Column(db.String(120), unique=True, nullable=False)
    ad = db.Column(db.String(300), nullable=False)
    aktif = db.Column(db.Boolean, default=True)

class AltProje(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    kod = db.Column(db.String(120), nullable=False)
    ad = db.Column(db.String(300), nullable=False)
    ana_proje_id = db.Column(db.Integer, db.ForeignKey('ana_proje.id'))
    ana_proje = db.relationship('AnaProje', backref='alt_projeler')
    aktif = db.Column(db.Boolean, default=True)

class FaturaProje(db.Model):
    """Bir faturaya bağlı çoklu proje satırları"""
    id = db.Column(db.Integer, primary_key=True)
    fatura_id = db.Column(db.Integer, db.ForeignKey('fatura.id'), nullable=False)
    ana_proje_id = db.Column(db.Integer, db.ForeignKey('ana_proje.id'))
    alt_proje_id = db.Column(db.Integer, db.ForeignKey('alt_proje.id'))
    sira = db.Column(db.Integer, default=1)
    ana_proje = db.relationship('AnaProje')

STANDART_ALT_KIRILIMLAR = (
    ('YI-EHS-ELKT', 'EHS Elektrik'),
    ('YI-EHS-MEK', 'EHS Mekanik'),
    ('YI-EHS-PNÖ', 'EHS Pnömatik'),
    ('YI-EHS-ÜSTYAPI', 'Üstyapı mekanik'),
    ('YI-EHS-SOFT', 'EHS Otomasyon Yazılım'),
    ('YI-FTS-ELKT', 'FTS Elektrik'),
    ('YI-FTS-MEK', 'FTS Mekanik'),
    ('YI-FTS-PNÖ', 'FTS Pnömatik'),
    ('YI-FTS-SOFT', 'FTS Otomasyon-Yazılım'),
    ('YI-PRJ-TŞRN', 'Taşeron tipi hizmetler'),
    ('YI-PRJ-BELGE', 'CE Analiz Sertifika'),
    ('YI-PRJ-HIRD', 'Hırdavat Boya'),
    ('YI-PRJ-İŞÇ-A', 'İşçilik-Atölye'),
    ('YI-PRJ-İŞÇ-S', 'İşçilik-Saha'),
    ('YI-PRJ-SEVK', 'Nakliye-Sigorta'),
    ('YI-PRJ-BG', 'Beklenmedik Gider'),
)

def standart_alt_kirilimlari_ekle(ana_projeler=None, commit=True):
    """Standart 16 alt kırılımı verilen veya bütün ana projelere tekilleştirerek ekle."""
    # Railway birden fazla worker başlattığında aynı başlangıç işleminin eş
    # zamanlı çalışmasını PostgreSQL işlem kilidiyle sıraya al.
    if db.engine.dialect.name == 'postgresql':
        db.session.execute(db.text('SELECT pg_advisory_xact_lock(2026080816)'))
    # Daha önce oluşmuş birebir tekrarları, fatura bağlantılarını koruyarak temizle.
    gorulen = {}
    for alt in AltProje.query.order_by(AltProje.id).all():
        anahtar = (alt.ana_proje_id, (alt.kod or '').strip().casefold(), (alt.ad or '').strip().casefold())
        if anahtar not in gorulen:
            gorulen[anahtar] = alt
            continue
        kalan = gorulen[anahtar]
        Fatura.query.filter_by(alt_proje_id=alt.id).update({'alt_proje_id': kalan.id}, synchronize_session=False)
        FaturaProje.query.filter_by(alt_proje_id=alt.id).update({'alt_proje_id': kalan.id}, synchronize_session=False)
        db.session.delete(alt)
    db.session.flush()
    projeler = list(ana_projeler) if ana_projeler is not None else AnaProje.query.all()
    eklenen = 0
    for ana in projeler:
        mevcut = {(a.kod or '').strip().casefold() for a in ana.alt_projeler}
        for kod, ad in STANDART_ALT_KIRILIMLAR:
            if kod.casefold() in mevcut:
                continue
            db.session.add(AltProje(kod=kod, ad=ad, ana_proje=ana))
            mevcut.add(kod.casefold())
            eklenen += 1
    db.session.flush()
    if db.engine.dialect.name == 'postgresql':
        db.session.execute(db.text(
            'CREATE UNIQUE INDEX IF NOT EXISTS uq_alt_proje_ana_kod_ad '
            'ON alt_proje (ana_proje_id, lower(kod), lower(ad))'
        ))
    if commit:
        db.session.commit()
    return eklenen
    alt_proje = db.relationship('AltProje')

class Fatura(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    dosya_adi = db.Column(db.String(255), nullable=False)
    fatura_no = db.Column(db.String(50))
    firma_adi = db.Column(db.String(200))
    fatura_tarihi = db.Column(db.String(20))
    tutar = db.Column(db.Float)
    para_birimi = db.Column(db.String(5), default='TRY')
    durum = db.Column(db.String(20), default='bekliyor')
    kaynak = db.Column(db.String(20), default='upload')
    yuklenme_tarihi = db.Column(db.DateTime, default=datetime.utcnow)
    not_alani = db.Column(db.Text)
    odeme_notu = db.Column(db.Text)
    onaylayan_id = db.Column(db.Integer, db.ForeignKey('kullanici.id'), nullable=True)
    onay_tarihi = db.Column(db.DateTime, nullable=True)
    departman_id = db.Column(db.Integer, db.ForeignKey('departman.id'), nullable=True)
    atanan_id = db.Column(db.Integer, db.ForeignKey('kullanici.id'), nullable=True)
    atama_tamamlandi = db.Column(db.Boolean, default=False, nullable=False)
    # Geriye dönük uyumluluk için tek proje alanları
    ana_proje_id = db.Column(db.Integer, db.ForeignKey('ana_proje.id'), nullable=True)
    alt_proje_id = db.Column(db.Integer, db.ForeignKey('alt_proje.id'), nullable=True)
    onaylayan = db.relationship('Kullanici', foreign_keys=[onaylayan_id])
    atanan = db.relationship('Kullanici', foreign_keys=[atanan_id])
    departman = db.relationship('Departman')
    ana_proje = db.relationship('AnaProje')
    alt_proje = db.relationship('AltProje')
    proje_satirlari = db.relationship('FaturaProje', backref='fatura', cascade='all, delete-orphan')

# ─── YARDIMCI ───────────────────────────────────────────────────────────────

def dosyadan_bilgi_cek(dosya_adi):
    isim = dosya_adi
    for ext in ('.pdf', '.PDF'):
        if isim.endswith(ext): isim = isim[:-len(ext)]; break
    if '_' in isim:
        idx = isim.index('_')
        return isim[:idx].strip(), isim[idx+1:].strip()
    return isim.strip(), ''

def metin_degeri(deger):
    """Kullanici metinlerini Turkce karakterleri bozmadan standart Unicode'a cevir."""
    if deger is None:
        return ''
    return unicodedata.normalize('NFC', str(deger)).strip()

def _arama_metni(deger):
    metin = unicodedata.normalize('NFKD', metin_degeri(deger)).casefold()
    metin = ''.join(ch for ch in metin if not unicodedata.combining(ch))
    metin = metin.replace('ı', 'i')
    return re.sub(r'[^a-z0-9]+', ' ', metin).strip()

def pdf_notundan_kullanici_bul(pdf_bytes):
    """PDF'deki Not alaninda adi gecen program kullanicisini bulur."""
    try:
        reader = PdfReader(_io.BytesIO(pdf_bytes))
        satirlar = []
        for page in reader.pages[:3]:
            satirlar.extend((page.extract_text() or '').splitlines())

        # E-fatura üreticilerinin metin katmanları farklıdır. PyPDF'in boş veya
        # parçalı döndürdüğü belgelerde PyMuPDF metnini de adaylara ekle.
        try:
            fitz_belge = fitz.open(stream=pdf_bytes, filetype='pdf')
            for sayfa_no in range(min(3, fitz_belge.page_count)):
                fitz_metni = fitz_belge.load_page(sayfa_no).get_text('text') or ''
                satirlar.extend(fitz_metni.splitlines())
            fitz_belge.close()
        except Exception as e:
            app.logger.debug('PyMuPDF metin okuma başarısız: %s', e)

        not_parcalari = []
        for i, satir in enumerate(satirlar):
            sade_satir = _arama_metni(satir)
            if re.search(r'(^| )not(?:lar)?( |$)', sade_satir):
                not_parcalari.extend(satirlar[i:i + 4])
        not_metni = f" {_arama_metni(' '.join(not_parcalari))} "
        if not not_metni.strip():
            return None

        kullanicilar = sorted(
            Kullanici.query.all(),
            key=lambda k: len(k.ad_soyad or ''),
            reverse=True,
        )
        for kullanici in kullanicilar:
            ad = _arama_metni(kullanici.ad_soyad)
            if len(ad.split()) >= 2 and f' {ad} ' in not_metni:
                return kullanici
    except Exception as e:
        app.logger.warning('PDF notundan kullanıcı bulunamadı: %s', e)
    return None

def metinden_kullanici_bul(metin):
    """Excel açıklamasındaki tam adı, yoksa benzersiz ad/ilk adı kullanıcıyla eşleştirir."""
    aranan = f" {_arama_metni(metin)} "
    if not aranan.strip():
        return None
    kullanicilar = sorted(Kullanici.query.all(), key=lambda k: len(k.ad_soyad or ''), reverse=True)
    for kullanici in kullanicilar:
        tam_ad = _arama_metni(kullanici.ad_soyad)
        if tam_ad and f' {tam_ad} ' in aranan:
            return kullanici
    # Açıklamada yalnızca "Özer Bey" gibi bir ilk ad bulunabilir. Aynı ilk
    # ada sahip birden fazla kullanıcı varsa yanlış atama yapmamak için geç.
    ilk_adlar = {}
    for kullanici in kullanicilar:
        parcalar = _arama_metni(kullanici.ad_soyad).split()
        if parcalar:
            ilk_adlar.setdefault(parcalar[0], []).append(kullanici)
    for ilk_ad, eslesenler in ilk_adlar.items():
        if len(eslesenler) == 1 and f' {ilk_ad} ' in aranan:
            return eslesenler[0]
    return None

def excel_fatura_satirlari(dosya_bytes, dosya_adi):
    """Eski XLS ve yeni XLSX dosyalarından ilk sayfanın satırlarını döndürür."""
    if dosya_adi.lower().endswith('.xls'):
        kitap = xlrd.open_workbook(file_contents=dosya_bytes)
        sayfa = kitap.sheet_by_index(0)
        return [[sayfa.cell_value(r, c) for c in range(sayfa.ncols)] for r in range(sayfa.nrows)]
    kitap = openpyxl.load_workbook(_io.BytesIO(dosya_bytes), read_only=True, data_only=True)
    sayfa = kitap.active
    return [list(satir) for satir in sayfa.iter_rows(values_only=True)]

def excel_sayi(deger):
    if deger is None or deger == '':
        return None
    if isinstance(deger, (int, float)):
        return float(deger)
    metin = str(deger).strip().replace(' ', '')
    if ',' in metin and '.' in metin:
        metin = metin.replace('.', '').replace(',', '.')
    elif ',' in metin:
        metin = metin.replace(',', '.')
    try:
        return float(metin)
    except ValueError:
        return None

def pdf_tarih_ve_tutar_bul(pdf_bytes):
    """PDF metnindeki fatura tarihi ve ödenecek/genel toplam tutarı bulur."""
    try:
        reader = PdfReader(_io.BytesIO(pdf_bytes))
        satirlar = []
        for page in reader.pages[:3]:
            satirlar.extend((page.extract_text() or '').splitlines())

        tarih = None
        tutar = None
        tarih_re = re.compile(
            r'(?<!\d)(?:(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})|'
            r'(\d{4})[./-](\d{1,2})[./-](\d{1,2}))(?!\d)'
        )
        tutar_re = re.compile(r'(?<!\d)(\d{1,3}(?:[.\s]\d{3})*,\d{2}|\d+[.,]\d{2})(?!\d)')

        def tarihi_duzenle(eslesme):
            if eslesme.group(1):
                gun, ay, yil = eslesme.group(1), eslesme.group(2), eslesme.group(3)
                if len(yil) == 2:
                    yil = '20' + yil
            else:
                yil, ay, gun = eslesme.group(4), eslesme.group(5), eslesme.group(6)
            try:
                return datetime(int(yil), int(ay), int(gun)).strftime('%d.%m.%Y')
            except ValueError:
                return None

        for i, satir in enumerate(satirlar):
            sade = _arama_metni(satir)
            baglam = ' '.join(satirlar[i:i + 3])
            if tarih is None and any(etiket in sade for etiket in (
                'fatura tarihi', 'fatura tarihi saati', 'invoice date',
                'duzenleme tarihi', 'belge tarihi', 'issue date'
            )):
                eslesme = tarih_re.search(baglam)
                if eslesme:
                    tarih = tarihi_duzenle(eslesme)
            if tutar is None and any(etiket in sade for etiket in (
                'odenecek tutar', 'genel toplam', 'vergiler dahil toplam',
                'payable amount', 'amount payable'
            )):
                eslesmeler = tutar_re.findall(baglam)
                if eslesmeler:
                    deger = eslesmeler[-1].replace(' ', '').replace('.', '').replace(',', '.')
                    tutar = float(deger)
            if tarih is not None and tutar is not None:
                break

        # Bazı e-fatura PDF'lerinde etiketler ile değerler farklı metin
        # bloklarında çıkar. Bu durumda bütün belgeyi tarayıp fatura tarihine
        # yakın olan geçerli tarihi seç; etiket yoksa ilk geçerli tarihi kullan.
        if tarih is None:
            tum_metin = '\n'.join(satirlar)
            adaylar = []
            for sira, eslesme in enumerate(tarih_re.finditer(tum_metin)):
                duzenli = tarihi_duzenle(eslesme)
                if not duzenli:
                    continue
                yil = int(duzenli[-4:])
                if yil < 2000 or yil > datetime.now().year + 1:
                    continue
                bas = max(0, eslesme.start() - 120)
                son = min(len(tum_metin), eslesme.end() + 120)
                cevre = _arama_metni(tum_metin[bas:son])
                puan = 0
                if any(x in cevre for x in (
                    'fatura tarihi', 'invoice date', 'duzenleme tarihi',
                    'belge tarihi', 'issue date'
                )):
                    puan += 20
                if any(x in cevre for x in (
                    'vade tarihi', 'odeme tarihi', 'due date',
                    'irsaliye tarihi', 'sevk tarihi'
                )):
                    puan -= 10
                adaylar.append((puan, -sira, duzenli))
            if adaylar:
                tarih = max(adaylar)[2]
        return tarih, tutar
    except Exception as e:
        app.logger.warning('PDF tarih/tutar bilgisi okunamadı: %s', e)
        return None, None

@app.route('/api/fatura-bilgileri-yeniden-oku', methods=['POST'])
@login_gerekli
def fatura_bilgileri_yeniden_oku():
    """Tarihi/tutarı eksik kayıtları saklanan PDF'lerden yeniden tamamla."""
    incelenen = guncellenen = 0
    for fatura in Fatura.query.all():
        if fatura.fatura_tarihi and fatura.tutar is not None:
            continue
        try:
            pdf_bytes, _ = fatura_pdf_bytes_ve_adres(fatura)
            tarih, tutar = pdf_tarih_ve_tutar_bul(pdf_bytes)
            incelenen += 1
            degisti = False
            if not fatura.fatura_tarihi and tarih:
                fatura.fatura_tarihi = tarih
                degisti = True
            if fatura.tutar is None and tutar is not None:
                fatura.tutar = tutar
                degisti = True
            if degisti:
                guncellenen += 1
        except Exception as e:
            app.logger.warning('Fatura bilgileri yeniden okunamadı (%s): %s', fatura.id, e)
    db.session.commit()
    return jsonify({'basarili': True, 'incelenen': incelenen, 'guncellenen': guncellenen})

def aktif_kullanici():
    kid = session.get('kullanici_id')
    return Kullanici.query.get(kid) if kid else None

def mim_fatura_silme_yetkisi():
    kullanici = aktif_kullanici()
    return bool(
        kullanici and kullanici.departman
        and metin_degeri(kullanici.departman.kod).casefold() == 'mim'
    )

def yonetici_mi():
    kullanici = aktif_kullanici()
    return bool(kullanici and kullanici.ad_soyad.strip().casefold() == 'dilek kaya')

def fatura_gorulebilir_mi(fatura):
    """Özel kullanıcıya atanan faturayı yalnızca o kullanıcı ve yönetici görür."""
    if yonetici_mi() or not fatura.atanan:
        return True
    if not fatura.atanan.atanan_faturalari_gizle:
        return True
    return fatura.atanan_id == session.get('kullanici_id')

def gorulebilir_fatura_veya_404(fatura_id):
    fatura = Fatura.query.get_or_404(fatura_id)
    if not fatura_gorulebilir_mi(fatura):
        abort(404)
    return fatura

def gorulebilir_fatura_sorgusu(sorgu=None):
    if sorgu is None:
        sorgu = Fatura.query
    if yonetici_mi():
        return sorgu
    kid = session.get('kullanici_id')
    return sorgu.filter(db.or_(
        Fatura.atanan_id.is_(None),
        Fatura.atanan_id == kid,
        ~Fatura.atanan.has(Kullanici.atanan_faturalari_gizle.is_(True)),
    ))

def supabase_dosya_yolu(public_url):
    """Supabase public URL'den dosya yolunu çıkar."""
    marker = f'/object/public/{SUPABASE_BUCKET}/'
    if marker in public_url:
        return public_url.split(marker)[-1]
    return public_url

def fatura_pdf_bytes_ve_adres(fatura):
    """Faturanın PDF içeriğini ve kullanıcıya gösterilecek dosya adını döndür."""
    kayit = (fatura.dosya_adi or '').strip()
    if not kayit:
        raise FileNotFoundError('Fatura dosya yolu boş')

    if kayit.startswith(('http://', 'https://')):
        netloc = urlparse(kayit).netloc.lower()
        if netloc == 'res.cloudinary.com':
            data = cloudinary_indir(kayit)
        elif SUPABASE_URL and SUPABASE_KEY:
            data = supabase_indir(kayit)
        else:
            r = http_requests.get(kayit, timeout=30)
            r.raise_for_status()
            data = r.content
        ad = os.path.basename(unquote(urlparse(kayit).path)) or 'fatura.pdf'
        return data, ad

    tam_yol = os.path.abspath(os.path.join(
        app.config['UPLOAD_FOLDER'], kayit.replace('/', os.sep)
    ))
    upload_root = os.path.abspath(app.config['UPLOAD_FOLDER'])
    if os.path.commonpath([tam_yol, upload_root]) != upload_root:
        raise FileNotFoundError('Geçersiz fatura dosya yolu')
    if not os.path.isfile(tam_yol):
        raise FileNotFoundError(tam_yol)
    with open(tam_yol, 'rb') as pdf_file:
        return pdf_file.read(), os.path.basename(tam_yol)

def fatura_onay_damgasi_uygula(fatura, kullanici=None):
    kullanici = kullanici or fatura.onaylayan or aktif_kullanici()
    proje_satirlari = sorted(fatura.proje_satirlari, key=lambda x: x.sira)
    projeler = [
        {
            'ana': satir.ana_proje.kod if satir.ana_proje else '-',
            'alt': satir.alt_proje.kod if satir.alt_proje else '-',
        }
        for satir in proje_satirlari
    ]
    if not projeler and (fatura.ana_proje or fatura.alt_proje):
        projeler = [{
            'ana': fatura.ana_proje.kod if fatura.ana_proje else '-',
            'alt': fatura.alt_proje.kod if fatura.alt_proje else '-',
        }]

    pdf_bytes, _ = fatura_pdf_bytes_ve_adres(fatura)
    damgali_pdf = pdf_damga_bytes(pdf_bytes, {
        'projeler': projeler,
        'not_alani': fatura.not_alani or '-',
        'odeme_notu': fatura.odeme_notu or '-',
        'onaylayan_ad': kullanici.ad_soyad if kullanici else '-',
        'departman_kod': (
            kullanici.departman.kod if kullanici and kullanici.departman
            else (fatura.departman.kod if fatura.departman else '-')
        ),
        'tarih': (fatura.onay_tarihi or datetime.utcnow()).strftime('%d.%m.%Y %H:%M'),
        'kisaltma': kullanici.kisaltma if kullanici else '',
    })

    if BULUT_MOD:
        if fatura.dosya_adi.startswith(SUPABASE_URL):
            eski_yol = supabase_dosya_yolu(fatura.dosya_adi)
            klasor = os.path.dirname(eski_yol).replace('\\', '/')
            kok_ad = os.path.splitext(os.path.basename(eski_yol))[0]
            # Aynı nesneye upsert edildiğinde CDN kısa süre eski PDF'yi
            # gösterebiliyor. Her damgada yeni URL oluşturarak bunu önle.
            yeni_ad = f"{secure_filename(kok_ad)}_onay_{datetime.now().strftime('%Y%m%d%H%M%S%f')}.pdf"
            dosya_yolu = f"{klasor}/{yeni_ad}" if klasor else yeni_ad
        else:
            guvenli_ad = secure_filename(
                f"{fatura.firma_adi or 'fatura'}_{fatura.fatura_no or fatura.id}.pdf"
            )
            kok_ad = os.path.splitext(guvenli_ad)[0]
            yeni_ad = f"{kok_ad}_onay_{datetime.now().strftime('%Y%m%d%H%M%S%f')}.pdf"
            dosya_yolu = f"faturalar/{datetime.now().year}/{datetime.now().month:02d}/{yeni_ad}"
        fatura.dosya_adi = supabase_yukle(damgali_pdf, dosya_yolu)
        db.session.commit()
    else:
        pdf_yolu = os.path.join(app.config['UPLOAD_FOLDER'],
                                fatura.dosya_adi.replace('/', os.sep))
        with open(pdf_yolu, 'wb') as damgali_dosya:
            damgali_dosya.write(damgali_pdf)

    return True

def seed_data():
    # Eski örnek kullanıcıyı kaldır; varsa onay ve atama geçmişini Dilek Kaya'ya taşı.
    eski_dilek = Kullanici.query.filter_by(ad_soyad='Dilek Yılmaz').first()
    dilek_kaya = Kullanici.query.filter_by(ad_soyad='Dilek Kaya').first()
    if eski_dilek:
        if dilek_kaya and dilek_kaya.id != eski_dilek.id:
            Fatura.query.filter_by(onaylayan_id=eski_dilek.id).update(
                {'onaylayan_id': dilek_kaya.id}, synchronize_session=False)
            Fatura.query.filter_by(atanan_id=eski_dilek.id).update(
                {'atanan_id': dilek_kaya.id}, synchronize_session=False)
            db.session.delete(eski_dilek)
        else:
            eski_dilek.ad_soyad = 'Dilek Kaya'
        db.session.commit()

    if Departman.query.count() == 0:
        db.session.add_all([
            Departman(kod='ASA', ad='Aftermarket Satış'),
            Departman(kod='ESA', ad='Endüstriyel Satış'),
            Departman(kod='SSA', ad='Servis'),
            Departman(kod='IDR', ad='İdari İşler'),
            Departman(kod='MUH', ad='Muhasebe'),
        ]); db.session.commit()
    if AnaProje.query.count() == 0:
        db.session.add_all([
            AnaProje(kod='GENEL', ad='Genel Giderler'),
            AnaProje(kod='ASA-2026', ad='Aftermarket 2026'),
            AnaProje(kod='ESA-2026', ad='Endüstriyel 2026'),
            AnaProje(kod='SSA-2026', ad='Servis 2026'),
        ]); db.session.commit()
    if AltProje.query.count() == 0:
        genel = AnaProje.query.filter_by(kod='GENEL').first()
        asa = AnaProje.query.filter_by(kod='ASA-2026').first()
        db.session.add_all([
            AltProje(kod='KIR', ad='Kırtasiye', ana_proje_id=genel.id),
            AltProje(kod='ULS', ad='Ulaşım', ana_proje_id=genel.id),
            AltProje(kod='PAZ', ad='Pazarlama', ana_proje_id=asa.id),
        ]); db.session.commit()
    if Kullanici.query.count() == 0:
        muh = Departman.query.filter_by(kod='MUH').first()
        db.session.add(Kullanici(ad_soyad='Dilek Kaya', departman_id=muh.id))
        db.session.commit()

# ─── LOGIN ──────────────────────────────────────────────────────────────────

@app.route('/login', methods=['GET','POST'])
def login():
    hata = None
    if request.method == 'POST':
        ad = metin_degeri(request.form.get('ad_soyad',''))
        sifre = request.form.get('sifre','')
        from unicodedata import normalize
        def normalize_tr(s):
            return s.replace('İ','i').replace('I','ı').replace('Ğ','g').replace('Ü','u') \
                     .replace('Ş','s').replace('Ö','o').replace('Ç','c') \
                     .replace('ı','i').replace('ğ','g').replace('ü','u') \
                     .replace('ş','s').replace('ö','o').replace('ç','c').lower()
        tum = Kullanici.query.all()
        k = next((x for x in tum if normalize_tr(x.ad_soyad) == normalize_tr(ad)), None)
        if not k:
            k = next((x for x in tum if normalize_tr(ad) in normalize_tr(x.ad_soyad)), None)
        if k and k.sifre_kontrol(sifre):
            session['kullanici_id'] = k.id
            session['kullanici_ad'] = k.ad_soyad
            session['kullanici_kisaltma'] = k.kisaltma
            session['departman_id'] = k.departman_id
            session['departman_kod'] = k.departman.kod if k.departman else ''
            return redirect(url_for('gelen_faturalar'))
        hata = 'Ad veya şifre hatalı'
    return render_template('login.html', hata=hata)

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/sifre-degistir', methods=['POST'])
@login_gerekli
def sifre_degistir():
    data = request.json
    k = aktif_kullanici()
    if not k: return jsonify({'hata': 'Oturum bulunamadı'}), 401
    yeni = data.get('yeni_sifre','').strip()
    if len(yeni) < 1: return jsonify({'hata': 'Şifre boş olamaz'}), 400
    k.sifre_ayarla(yeni)
    db.session.commit()
    return jsonify({'basarili': True})

# ─── ANA SAYFA ──────────────────────────────────────────────────────────────

@app.route('/')
@login_gerekli
def index():
    return redirect(url_for('gelen_faturalar'))

@app.route('/gelen')
@login_gerekli
def gelen_faturalar():
    durum = request.args.get('durum', 'bekliyor')
    kid = session.get('kullanici_id')
    filtre = request.args.get('filtre', 'tum')

    q = gorulebilir_fatura_sorgusu()
    if durum != 'tum':
        q = q.filter_by(durum=durum)
    if filtre == 'bana':
        q = q.filter_by(atanan_id=kid)

    faturalar = q.order_by(Fatura.yuklenme_tarihi.desc()).all()
    kullanicilar = Kullanici.query.all()
    ana_projeler = AnaProje.query.filter_by(aktif=True).all()

    bana_count = gorulebilir_fatura_sorgusu().filter_by(durum='bekliyor', atanan_id=kid).count()
    sayim_q = gorulebilir_fatura_sorgusu()
    durum_sayilari = {
        'tum': sayim_q.count(),
        'bekliyor': sayim_q.filter_by(durum='bekliyor').count(),
        'onaylandi': sayim_q.filter_by(durum='onaylandi').count(),
        'reddedildi': sayim_q.filter_by(durum='reddedildi').count(),
        'eksik_bilgi': sayim_q.filter_by(durum='eksik_bilgi').count(),
    }
    return render_template('gelen.html',
        faturalar=faturalar, kullanicilar=kullanicilar,
        ana_projeler=ana_projeler, aktif_durum=durum,
        durum_sayilari=durum_sayilari, filtre=filtre,
        bana_count=bana_count,
        oturum_kullanici=aktif_kullanici()
    )

# ─── FATURA API ─────────────────────────────────────────────────────────────

@app.route('/api/upload', methods=['POST'])
@login_gerekli
def upload_fatura():
    if 'dosyalar' not in request.files:
        return jsonify({'hata': 'Dosya bulunamadı'}), 400
    yuklenenler = []
    for dosya in request.files.getlist('dosyalar'):
        if dosya.filename.lower().endswith('.pdf'):
            dosya_adi = dosya.filename
            dosya_bytes = dosya.read()
            not_kullanicisi = pdf_notundan_kullanici_bul(dosya_bytes)
            pdf_tarihi, pdf_tutari = pdf_tarih_ve_tutar_bul(dosya_bytes)
            simdi = datetime.now()
            if BULUT_MOD:
                # Supabase Storage'a yükle
                depo_dosya_adi = secure_filename(dosya_adi)
                if not depo_dosya_adi.lower().endswith('.pdf'):
                    depo_dosya_adi = f"fatura_{simdi.strftime('%Y%m%d%H%M%S%f')}.pdf"
                dosya_yolu = f"faturalar/{simdi.year}/{simdi.month:02d}/{depo_dosya_adi}"
                goreli_yol = supabase_yukle(dosya_bytes, dosya_yolu)
            else:
                # Lokal - aylık klasör
                ay_klasor = os.path.join(app.config['UPLOAD_FOLDER'],
                                         str(simdi.year), f'{simdi.month:02d}')
                os.makedirs(ay_klasor, exist_ok=True)
                with open(os.path.join(ay_klasor, dosya_adi), 'wb') as yerel_pdf:
                    yerel_pdf.write(dosya_bytes)
                goreli_yol = f"{simdi.year}/{simdi.month:02d}/{dosya_adi}"
            fno, firma = dosyadan_bilgi_cek(dosya_adi)
            sade_dosya_adi = secure_filename(dosya_adi).casefold()
            mevcut = next(
                (
                    kayit for kayit in Fatura.query.all()
                    if secure_filename(
                        os.path.basename(unquote(kayit.dosya_adi or ''))
                    ).casefold() == sade_dosya_adi
                    or (fno and kayit.fatura_no == fno)
                ),
                None,
            )
            if mevcut:
                # Aynı faturanın eski/kırık bulut bağlantısını yeni Supabase
                # adresiyle değiştir; ikinci bir fatura kaydı oluşturma.
                mevcut.dosya_adi = goreli_yol
                # Bekleyen bir fatura yeniden yüklendiyse Toplu Atama listesinde
                # tekrar görünsün. Onaylanmış/reddedilmiş durumunu değiştirme.
                if mevcut.durum == 'bekliyor':
                    mevcut.atama_tamamlandi = False
                if pdf_tarihi:
                    mevcut.fatura_tarihi = pdf_tarihi
                if pdf_tutari is not None:
                    mevcut.tutar = pdf_tutari
                if not_kullanicisi:
                    mevcut.atanan_id = not_kullanicisi.id
                yuklenenler.append(dosya_adi)
            elif not Fatura.query.filter_by(dosya_adi=goreli_yol).first():
                db.session.add(Fatura(dosya_adi=goreli_yol, kaynak='upload',
                                      fatura_no=fno or None, firma_adi=firma or None,
                                      fatura_tarihi=pdf_tarihi, tutar=pdf_tutari,
                                      atanan_id=not_kullanicisi.id if not_kullanicisi else None))
                yuklenenler.append(dosya_adi)
    db.session.commit()
    return jsonify({'basarili': True, 'yuklenen': len(yuklenenler)})

@app.route('/api/excel-fatura-bilgileri', methods=['POST'])
@login_gerekli
def excel_fatura_bilgileri():
    dosya = request.files.get('dosya')
    if not dosya or not dosya.filename.lower().endswith(('.xls', '.xlsx')):
        return jsonify({'hata': 'Lütfen .xls veya .xlsx dosyası seçin'}), 400
    try:
        satirlar = excel_fatura_satirlari(dosya.read(), dosya.filename)
        baslik_satiri = None
        kolonlar = {}
        for satir_no, satir in enumerate(satirlar[:15]):
            sade = [_arama_metni(x) for x in satir]
            for i, baslik in enumerate(sade):
                if baslik == 'fatura no': kolonlar['fatura_no'] = i
                elif baslik == 'fatura tutari': kolonlar['tutar'] = i
                elif baslik == 'fatura aciklamasi': kolonlar['aciklama'] = i
            if 'fatura_no' in kolonlar and 'tutar' in kolonlar:
                baslik_satiri = satir_no
                break
        if baslik_satiri is None:
            return jsonify({'hata': 'Fatura No ve Fatura Tutarı sütunları bulunamadı'}), 400

        faturalar = {
            re.sub(r'[^a-z0-9]', '', _arama_metni(f.fatura_no)): f
            for f in Fatura.query.all() if f.fatura_no
        }
        eslesen = tutar_guncellenen = atanan = 0
        eslesmeyen = []
        for satir in satirlar[baslik_satiri + 1:]:
            if kolonlar['fatura_no'] >= len(satir):
                continue
            ham_no = satir[kolonlar['fatura_no']]
            if ham_no is None or ham_no == '':
                continue
            fatura_no = str(int(ham_no)) if isinstance(ham_no, float) and ham_no.is_integer() else str(ham_no).strip()
            anahtar = re.sub(r'[^a-z0-9]', '', _arama_metni(fatura_no))
            fatura = faturalar.get(anahtar)
            if not fatura:
                eslesmeyen.append(fatura_no)
                continue
            eslesen += 1
            tutar = excel_sayi(satir[kolonlar['tutar']] if kolonlar['tutar'] < len(satir) else None)
            if tutar is not None:
                fatura.tutar = tutar
                tutar_guncellenen += 1
            if 'aciklama' in kolonlar and kolonlar['aciklama'] < len(satir):
                kullanici = metinden_kullanici_bul(satir[kolonlar['aciklama']])
                if kullanici:
                    fatura.atanan_id = kullanici.id
                    atanan += 1
        db.session.commit()
        return jsonify({
            'basarili': True, 'eslesen': eslesen,
            'tutar_guncellenen': tutar_guncellenen, 'atanan': atanan,
            'eslesmeyen': len(eslesmeyen)
        })
    except Exception as e:
        app.logger.exception('Excel fatura bilgileri yüklenemedi')
        return jsonify({'hata': f'Excel okunamadı: {e}'}), 400

@app.route('/api/network-tara', methods=['POST'])
@login_gerekli
def network_tara():
    klasor = request.json.get('klasor','')
    if not klasor or not os.path.exists(klasor):
        return jsonify({'hata': 'Klasör bulunamadı'}), 400
    eklenen = 0
    for dosya_adi in os.listdir(klasor):
        if dosya_adi.lower().endswith('.pdf'):
            simdi = datetime.now()
            ay_klasor = os.path.join(app.config['UPLOAD_FOLDER'],
                                     str(simdi.year), f'{simdi.month:02d}')
            os.makedirs(ay_klasor, exist_ok=True)
            hedef = os.path.join(ay_klasor, dosya_adi)
            if not os.path.exists(hedef):
                shutil.copy2(os.path.join(klasor, dosya_adi), hedef)
            goreli_yol = f"{simdi.year}/{simdi.month:02d}/{dosya_adi}"
            if not Fatura.query.filter_by(dosya_adi=goreli_yol).first():
                fno, firma = dosyadan_bilgi_cek(dosya_adi)
                db.session.add(Fatura(dosya_adi=goreli_yol, kaynak='network',
                                      fatura_no=fno or None, firma_adi=firma or None))
                eklenen += 1
    db.session.commit()
    return jsonify({'basarili': True, 'eklenen': eklenen})

@app.route('/api/fatura/<int:fatura_id>/detay')
@login_gerekli
def fatura_detay(fatura_id):
    f = gorulebilir_fatura_veya_404(fatura_id)
    projeler = [{'ana_proje_id': p.ana_proje_id, 'alt_proje_id': p.alt_proje_id,
                 'ana_kod': p.ana_proje.kod if p.ana_proje else '',
                 'ana_ad': p.ana_proje.ad if p.ana_proje else '',
                 'alt_kod': p.alt_proje.kod if p.alt_proje else '',
                 'alt_ad': p.alt_proje.ad if p.alt_proje else ''}
                for p in sorted(f.proje_satirlari, key=lambda x: x.sira)]
    pdf_url = f.dosya_adi
    return jsonify({
        'id': f.id, 'fatura_no': f.fatura_no or '',
        'firma_adi': f.firma_adi or '', 'fatura_tarihi': f.fatura_tarihi or '',
        'tutar': f.tutar or '', 'not_alani': f.not_alani or '',
        'odeme_notu': f.odeme_notu or '',
        'projeler': projeler,
        'onaylayan_id': f.onaylayan_id, 'atanan_id': f.atanan_id,
        'pdf_url': pdf_url,
    })

@app.route('/api/fatura/<int:fatura_id>/bilgi', methods=['POST'])
@login_gerekli
def fatura_bilgi_guncelle(fatura_id):
    f = gorulebilir_fatura_veya_404(fatura_id)
    d = request.json
    f.fatura_no = d.get('fatura_no', f.fatura_no)
    f.firma_adi = metin_degeri(d.get('firma_adi', f.firma_adi))
    f.fatura_tarihi = d.get('fatura_tarihi', f.fatura_tarihi)
    f.tutar = d.get('tutar', f.tutar)
    db.session.commit()
    return jsonify({'basarili': True})

@app.route('/api/fatura/<int:fatura_id>', methods=['DELETE'])
@login_gerekli
def fatura_sil(fatura_id):
    if not mim_fatura_silme_yetkisi():
        return jsonify({'hata': 'Fatura silme yetkisi yalnızca MIM departmanına aittir.'}), 403

    fatura = gorulebilir_fatura_veya_404(fatura_id)
    db.session.delete(fatura)
    db.session.commit()
    return jsonify({'basarili': True})

@app.route('/api/fatura/<int:fatura_id>/damgala', methods=['POST'])
@login_gerekli
def onayli_faturayi_damgala(fatura_id):
    fatura = gorulebilir_fatura_veya_404(fatura_id)
    if fatura.durum != 'onaylandi':
        return jsonify({'hata': 'Yalnızca onaylanmış faturalar damgalanabilir.'}), 400
    try:
        fatura_onay_damgasi_uygula(fatura)
        return jsonify({'basarili': True})
    except Exception as e:
        app.logger.exception('Onay damgası hatası: %s', e)
        return jsonify({'hata': 'Onay damgası PDF dosyasına eklenemedi.'}), 500

@app.route('/api/fatura/<int:fatura_id>/onayla', methods=['POST'])
@login_gerekli
def fatura_onayla(fatura_id):
    fatura = gorulebilir_fatura_veya_404(fatura_id)
    data = request.json

    # Çoklu proje satırlarını kaydet
    FaturaProje.query.filter_by(fatura_id=fatura_id).delete()
    projeler = data.get('projeler', [])
    for i, p in enumerate(projeler):
        if p.get('ana_proje_id'):
            db.session.add(FaturaProje(
                fatura_id=fatura_id,
                ana_proje_id=p['ana_proje_id'] or None,
                alt_proje_id=p.get('alt_proje_id') or None,
                sira=i+1
            ))
    if projeler and projeler[0].get('ana_proje_id'):
        fatura.ana_proje_id = projeler[0]['ana_proje_id']
        fatura.alt_proje_id = projeler[0].get('alt_proje_id')

    # Fatura bilgilerini onayla aynı işlemde kaydet. Böylece ayrı bilgi
    # isteği aksasa bile Onaylananlar ekranında firma adı boş kalmaz.
    fatura.fatura_no = metin_degeri(data.get('fatura_no', fatura.fatura_no)) or fatura.fatura_no
    fatura.firma_adi = metin_degeri(data.get('firma_adi', fatura.firma_adi)) or fatura.firma_adi
    fatura.fatura_tarihi = metin_degeri(data.get('fatura_tarihi', fatura.fatura_tarihi)) or fatura.fatura_tarihi
    if data.get('tutar') not in (None, ''):
        fatura.tutar = data.get('tutar')
    fatura.not_alani = metin_degeri(data.get('not_alani',''))
    fatura.odeme_notu = metin_degeri(data.get('odeme_notu',''))
    fatura.onaylayan_id = data.get('kullanici_id') or session.get('kullanici_id')
    fatura.departman_id = data.get('departman_id') or session.get('departman_id')
    fatura.durum = data.get('durum', 'onaylandi')
    fatura.onay_tarihi = datetime.utcnow()
    db.session.commit()

    # Onaylananlar için PDF damga
    damga_uygulandi = None
    if fatura.durum == 'onaylandi':
        try:
            fatura_onay_damgasi_uygula(fatura, aktif_kullanici())
            damga_uygulandi = True
        except Exception as e:
            damga_uygulandi = False
            app.logger.exception('Damga hatası: %s', e)

    return jsonify({'basarili': True, 'damga_uygulandi': damga_uygulandi})

@app.route('/api/fatura/<int:fatura_id>/aktar', methods=['POST'])
@login_gerekli
def fatura_aktar(fatura_id):
    fatura = gorulebilir_fatura_veya_404(fatura_id)
    data = request.json
    hedef_id = data.get('hedef_kullanici_id')
    if not hedef_id:
        return jsonify({'hata': 'Hedef kullanıcı seçin'}), 400
    fatura.atanan_id = hedef_id
    fatura.durum = 'bekliyor'
    db.session.commit()
    hedef = Kullanici.query.get(hedef_id)
    return jsonify({'basarili': True, 'hedef': hedef.ad_soyad if hedef else ''})

@app.route('/toplu-atama')
@login_gerekli
def toplu_atama_sayfasi():
    if not mim_fatura_silme_yetkisi():
        return jsonify({'hata': 'Bu ekran yalnızca MIM departmanına açıktır.'}), 403
    faturalar = Fatura.query.filter_by(durum='bekliyor', atama_tamamlandi=False).order_by(
        Fatura.yuklenme_tarihi.desc()).all()
    kullanicilar = Kullanici.query.order_by(Kullanici.ad_soyad).all()
    return render_template('toplu_atama.html', faturalar=faturalar,
                           kullanicilar=kullanicilar,
                           oturum_kullanici=aktif_kullanici())

@app.route('/api/toplu-atama', methods=['POST'])
@login_gerekli
def api_toplu_atama():
    if not mim_fatura_silme_yetkisi():
        return jsonify({'hata': 'Toplu atama yetkisi yalnızca MIM departmanına aittir.'}), 403
    data = request.json or {}
    fatura_idleri = data.get('fatura_idleri') or []
    kullanici_id = data.get('kullanici_id')
    if not fatura_idleri or not kullanici_id:
        return jsonify({'hata': 'Fatura ve kullanıcı seçimi gereklidir.'}), 400
    kullanici = Kullanici.query.get_or_404(kullanici_id)
    faturalar = Fatura.query.filter(Fatura.id.in_(fatura_idleri)).all()
    for fatura in faturalar:
        fatura.atanan_id = kullanici.id
        fatura.durum = 'bekliyor'
    db.session.commit()
    return jsonify({'basarili': True, 'atanan': len(faturalar),
                    'kullanici': kullanici.ad_soyad})

@app.route('/api/toplu-atama/tamamla', methods=['POST'])
@login_gerekli
def api_toplu_atama_tamamla():
    if not mim_fatura_silme_yetkisi():
        return jsonify({'hata': 'Bu işlem yalnızca MIM departmanına aittir.'}), 403
    fatura_idleri = (request.json or {}).get('fatura_idleri') or []
    if not fatura_idleri:
        return jsonify({'hata': 'En az bir fatura seçin.'}), 400
    faturalar = Fatura.query.filter(Fatura.id.in_(fatura_idleri)).all()
    for fatura in faturalar:
        fatura.atama_tamamlandi = True
    db.session.commit()
    return jsonify({'basarili': True, 'tamamlanan': len(faturalar)})

@app.route('/api/alt-projeler/<int:ana_proje_id>')
@login_gerekli
def alt_projeler(ana_proje_id):
    altlar = AltProje.query.filter_by(ana_proje_id=ana_proje_id, aktif=True).all()
    return jsonify([{'id': a.id, 'kod': a.kod, 'ad': a.ad} for a in altlar])

@app.route('/uploads/<path:dosya_adi>')
@login_gerekli
def serve_pdf(dosya_adi):
    if dosya_adi.startswith('http'):
        # Supabase public URL'den proxy olarak çek
        try:
            data = supabase_indir(dosya_adi)
            from flask import Response
            return Response(data, mimetype='application/pdf',
                            headers={'Content-Disposition': 'inline; filename="fatura.pdf"'})
        except Exception as e:
            print(f"PDF proxy hatasi: {e}")
            return redirect(dosya_adi)
    return send_from_directory(app.config['UPLOAD_FOLDER'], dosya_adi.replace('/', os.sep))

@app.route('/api/fatura/<int:fatura_id>/pdf')
@login_gerekli
def fatura_pdf(fatura_id):
    """PDF'yi veritabanındaki kayıt üzerinden güvenli biçimde tarayıcıya gönder."""
    fatura = gorulebilir_fatura_veya_404(fatura_id)
    try:
        data, dosya_adi = fatura_pdf_bytes_ve_adres(fatura)
        dosya_adi = secure_filename(dosya_adi) or 'fatura.pdf'
        if not data.startswith(b'%PDF'):
            raise ValueError('Depolama servisi geçerli bir PDF döndürmedi')
        return Response(
            data,
            mimetype='application/pdf',
            headers={
                'Content-Disposition': f'inline; filename="{dosya_adi}"',
                'Cache-Control': 'private, no-store',
                'X-Content-Type-Options': 'nosniff',
            },
        )
    except FileNotFoundError:
        return jsonify({
            'hata': 'PDF dosyası bulunamadı. Railway geçici diskinden silinmiş olabilir.'
        }), 404
    except Exception as e:
        app.logger.exception('Fatura PDF görüntüleme hatası: %s', e)
        return jsonify({
            'hata': 'PDF depolama servisinden alınamadı. Supabase ayarlarını kontrol edin.'
        }), 502

@app.route('/fatura/<int:fatura_id>/goruntule')
@login_gerekli
def fatura_goruntule(fatura_id):
    """PDF'yi tarayıcı eklentisine gerek kalmadan resim sayfaları olarak göster."""
    fatura = gorulebilir_fatura_veya_404(fatura_id)
    try:
        import fitz
        data, dosya_adi = fatura_pdf_bytes_ve_adres(fatura)
        if not data.startswith(b'%PDF'):
            raise ValueError('Geçerli bir PDF dosyası değil')

        # Onay sırasında depolama kısa süreli hata verirse damga eksik kalmasın.
        # Kullanıcı faturayı görüntülediğinde eksik damga bir kez otomatik tamamlanır.
        if fatura.durum == 'onaylandi' and not pdf_damgali_mi(data):
            try:
                fatura_onay_damgasi_uygula(fatura)
                data, dosya_adi = fatura_pdf_bytes_ve_adres(fatura)
            except Exception as damga_hatasi:
                app.logger.exception('Görüntülemede otomatik damga hatası: %s', damga_hatasi)

        belge = fitz.open(stream=data, filetype='pdf')
        sayfalar = []
        for sayfa in belge:
            goruntu = sayfa.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
            png = goruntu.tobytes('png')
            sayfalar.append(base64.b64encode(png).decode('ascii'))
        belge.close()

        return render_template('pdf_goruntule.html', fatura=fatura,
                               dosya_adi=dosya_adi, sayfalar=sayfalar)
    except Exception as e:
        app.logger.exception('Fatura resimli görüntüleme hatası: %s', e)
        return render_template('pdf_goruntule.html', fatura=fatura,
                               dosya_adi='Fatura', sayfalar=[],
                               hata='Fatura görüntülenemedi. Lütfen tekrar deneyin.'), 502

# ─── MUHASEBE ───────────────────────────────────────────────────────────────

@app.route('/muhasebe')
@login_gerekli
def muhasebe_panel():
    faturalar = (gorulebilir_fatura_sorgusu()
                 .filter_by(durum='onaylandi')
                 .order_by(Fatura.onay_tarihi.desc()).all())
    return render_template('muhasebe.html', faturalar=faturalar, now=datetime.utcnow(),
                           oturum_kullanici=aktif_kullanici())

@app.route('/api/river-export')
@login_gerekli
def river_export():
    from openpyxl import Workbook
    from flask import make_response
    import io
    faturalar = gorulebilir_fatura_sorgusu().filter_by(durum='onaylandi').all()
    wb = Workbook(); ws = wb.active; ws.title = "River Sipariş"
    ws.append(['Fatura No','Tarih','Firma','Tutar','Para Birimi',
               'Ana Proje','Alt Proje','Not','Ödeme Notu','Onaylayan','Departman','Onay Tarihi'])
    for f in faturalar:
        projeler = sorted(f.proje_satirlari, key=lambda x: x.sira)
        ana = ', '.join([p.ana_proje.kod for p in projeler if p.ana_proje]) or (f.ana_proje.kod if f.ana_proje else '')
        alt = ', '.join([p.alt_proje.kod for p in projeler if p.alt_proje]) or (f.alt_proje.kod if f.alt_proje else '')
        ws.append([f.fatura_no or '', f.fatura_tarihi or '', f.firma_adi or '',
                   f.tutar or '', f.para_birimi or 'TRY', ana, alt,
                   f.not_alani or '', f.odeme_notu or '',
                   f'{f.onaylayan.ad_soyad} ({f.onaylayan.kisaltma})' if f.onaylayan else '',
                   f.departman.kod if f.departman else '',
                   f.onay_tarihi.strftime('%d.%m.%Y %H:%M') if f.onay_tarihi else ''])
    out = io.BytesIO(); wb.save(out); out.seek(0)
    resp = make_response(out.read())
    resp.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    resp.headers['Content-Disposition'] = 'attachment; filename=river_siparis.xlsx'
    return resp

# ─── YÖNETİM ────────────────────────────────────────────────────────────────

@app.route('/yonetim/kullanicilar')
@app.route('/yonetim/projeler')
@login_gerekli
def yonetim():
    return render_template('yonetim.html', oturum_kullanici=aktif_kullanici())

@app.route('/api/yonetim/kullanicilar')
@login_gerekli
def api_kullanicilar():
    return jsonify([{'id':k.id,'ad_soyad':k.ad_soyad,
                     'email':k.email or '',
                     'atanan_faturalari_gizle': bool(k.atanan_faturalari_gizle),
                     'gizlilik_duzenlenebilir': yonetici_mi(),
                     'departman_id':k.departman_id,
                     'departman_kod':k.departman.kod if k.departman else None,
                     'departman_ad':k.departman.ad if k.departman else None}
                    for k in Kullanici.query.order_by(Kullanici.ad_soyad).all()])

@app.route('/api/yonetim/kullanici-ekle', methods=['POST'])
@login_gerekli
def api_kullanici_ekle():
    d = request.json; ad = metin_degeri(d.get('ad_soyad',''))
    if not ad: return jsonify({'hata':'Ad Soyad boş olamaz'}),400
    def norm(s):
        return s.replace('İ','i').replace('I','ı').replace('Ğ','g').replace('Ü','u') \
               .replace('Ş','s').replace('Ö','o').replace('Ç','c') \
               .replace('ı','i').replace('ğ','g').replace('ü','u') \
               .replace('ş','s').replace('ö','o').replace('ç','c').lower().strip()
    mevcut = next((k for k in Kullanici.query.all() if norm(k.ad_soyad) == norm(ad)), None)
    if mevcut:
        return jsonify({'hata': f'"{mevcut.ad_soyad}" isimli kullanıcı zaten var'}), 400
    email = metin_degeri(d.get('email', '')).lower()
    if email and ('@' not in email or '.' not in email.rsplit('@', 1)[-1]):
        return jsonify({'hata':'Geçerli bir e-posta adresi girin'}), 400
    if email and Kullanici.query.filter(db.func.lower(Kullanici.email) == email).first():
        return jsonify({'hata':'Bu e-posta adresi başka bir kullanıcıda kayıtlı'}), 400
    k = Kullanici(ad_soyad=ad, email=email or None,
                  departman_id=d.get('departman_id') or None)
    db.session.add(k); db.session.commit()
    return jsonify({'basarili':True,'id':k.id})

@app.route('/api/yonetim/kullanici-guncelle/<int:kid>', methods=['POST'])
@login_gerekli
def api_kullanici_guncelle(kid):
    k = Kullanici.query.get_or_404(kid)
    d = request.json
    ad = metin_degeri(d.get('ad_soyad',''))
    if not ad: return jsonify({'hata':'Ad Soyad boş olamaz'}),400
    def norm(s):
        return s.replace('İ','i').replace('I','ı').replace('Ğ','g').replace('Ü','u') \
               .replace('Ş','s').replace('Ö','o').replace('Ç','c') \
               .replace('ı','i').replace('ğ','g').replace('ü','u') \
               .replace('ş','s').replace('ö','o').replace('ç','c').lower().strip()
    mevcut = next((x for x in Kullanici.query.all() if norm(x.ad_soyad)==norm(ad) and x.id!=kid), None)
    if mevcut:
        return jsonify({'hata': f'"{mevcut.ad_soyad}" isimli kullanıcı zaten var'}), 400
    k.ad_soyad = ad
    email = metin_degeri(d.get('email', '')).lower()
    if email and ('@' not in email or '.' not in email.rsplit('@', 1)[-1]):
        return jsonify({'hata':'Geçerli bir e-posta adresi girin'}), 400
    email_sahibi = (Kullanici.query
                    .filter(db.func.lower(Kullanici.email) == email,
                            Kullanici.id != kid).first()) if email else None
    if email_sahibi:
        return jsonify({'hata':'Bu e-posta adresi başka bir kullanıcıda kayıtlı'}), 400
    k.email = email or None
    k.departman_id = d.get('departman_id') or None
    if yonetici_mi() and 'atanan_faturalari_gizle' in d:
        k.atanan_faturalari_gizle = bool(d.get('atanan_faturalari_gizle'))
    db.session.commit()
    return jsonify({'basarili':True})

@app.route('/api/yonetim/kullanici-sil/<int:kid>', methods=['DELETE'])
@login_gerekli
def api_kullanici_sil(kid):
    db.session.delete(Kullanici.query.get_or_404(kid)); db.session.commit()
    return jsonify({'basarili':True})

@app.route('/api/yonetim/departmanlar')
@login_gerekli
def api_departmanlar():
    return jsonify([{'id':d.id,'kod':d.kod,'ad':d.ad} for d in Departman.query.all()])

@app.route('/api/yonetim/departman-ekle', methods=['POST'])
@login_gerekli
def api_departman_ekle():
    d = request.json; kod = metin_degeri(d.get('kod','')).upper(); ad = metin_degeri(d.get('ad',''))
    if not kod or not ad: return jsonify({'hata':'Kod ve Ad boş olamaz'}),400
    if Departman.query.filter_by(kod=kod).first(): return jsonify({'hata':f'"{kod}" zaten var'}),400
    dep = Departman(kod=kod,ad=ad); db.session.add(dep); db.session.commit()
    return jsonify({'basarili':True,'id':dep.id})

@app.route('/api/yonetim/departman-guncelle/<int:did>', methods=['POST'])
@login_gerekli
def api_departman_guncelle(did):
    d = Departman.query.get_or_404(did)
    data = request.json
    kod = metin_degeri(data.get('kod','')).upper()
    ad = metin_degeri(data.get('ad',''))
    if not kod or not ad: return jsonify({'hata':'Kod ve Ad boş olamaz'}),400
    mevcut = Departman.query.filter_by(kod=kod).first()
    if mevcut and mevcut.id != did: return jsonify({'hata':f'"{kod}" kodu zaten var'}),400
    d.kod = kod; d.ad = ad
    db.session.commit()
    return jsonify({'basarili':True})

@app.route('/api/yonetim/departman-sil/<int:did>', methods=['DELETE'])
@login_gerekli
def api_departman_sil(did):
    db.session.delete(Departman.query.get_or_404(did)); db.session.commit()
    return jsonify({'basarili':True})

@app.route('/api/yonetim/projeler')
@login_gerekli
def api_projeler():
    return jsonify([{'id':p.id,'kod':p.kod,'ad':p.ad,'aktif':p.aktif,
                     'silinebilir':yonetici_mi(),
                     'alt_projeler':[{'id':a.id,'kod':a.kod,'ad':a.ad,'aktif':a.aktif}
                                     for a in p.alt_projeler]}
                    for p in AnaProje.query.all()])

@app.route('/api/yonetim/ana-proje-ekle', methods=['POST'])
@login_gerekli
def api_ana_proje_ekle():
    d = request.json; kod = metin_degeri(d.get('kod','')).upper(); ad = metin_degeri(d.get('ad',''))
    if not kod or not ad: return jsonify({'hata':'Eksik bilgi'}),400
    if AnaProje.query.filter_by(kod=kod).first(): return jsonify({'hata':f'"{kod}" zaten var'}),400
    p = AnaProje(kod=kod,ad=ad); db.session.add(p); db.session.flush()
    standart_alt_kirilimlari_ekle([p], commit=False)
    db.session.commit()
    return jsonify({'basarili':True,'id':p.id})

@app.route('/api/yonetim/ana-proje-sil/<int:pid>', methods=['DELETE'])
@login_gerekli
def api_ana_proje_sil(pid):
    if not yonetici_mi(): return jsonify({'hata':'Bu işlem yalnızca yönetici Dilek Kaya tarafından yapılabilir.'}),403
    p = AnaProje.query.get_or_404(pid)
    if Fatura.query.filter_by(ana_proje_id=pid).first() or FaturaProje.query.filter_by(ana_proje_id=pid).first():
        return jsonify({'hata':'Bu proje faturalarda kullanıldığı için silinemez; pasife alabilirsiniz.'}),400
    for alt in list(p.alt_projeler):
        if Fatura.query.filter_by(alt_proje_id=alt.id).first() or FaturaProje.query.filter_by(alt_proje_id=alt.id).first():
            return jsonify({'hata':'Bu projenin kullanılan alt projeleri var; pasife alabilirsiniz.'}),400
    for alt in list(p.alt_projeler): db.session.delete(alt)
    db.session.delete(p); db.session.commit()
    return jsonify({'basarili':True})

@app.route('/api/yonetim/ana-proje-toggle/<int:pid>', methods=['POST'])
@login_gerekli
def api_ana_proje_toggle(pid):
    p = AnaProje.query.get_or_404(pid); p.aktif = request.json.get('aktif', not p.aktif)
    db.session.commit(); return jsonify({'basarili':True})

@app.route('/api/yonetim/alt-proje-ekle', methods=['POST'])
@login_gerekli
def api_alt_proje_ekle():
    d = request.json; kod = metin_degeri(d.get('kod','')).upper(); ad = metin_degeri(d.get('ad',''))
    if not kod or not ad or not d.get('ana_proje_id'): return jsonify({'hata':'Eksik bilgi'}),400
    if AltProje.query.filter_by(kod=kod,ad=ad,ana_proje_id=d['ana_proje_id']).first():
        return jsonify({'hata':f'"{kod} — {ad}" alt kırılımı zaten var'}),400
    a = AltProje(kod=kod,ad=ad,ana_proje_id=d['ana_proje_id']); db.session.add(a); db.session.commit()
    return jsonify({'basarili':True,'id':a.id})

@app.route('/api/yonetim/alt-proje-sil/<int:aid>', methods=['DELETE'])
@login_gerekli
def api_alt_proje_sil(aid):
    if not yonetici_mi(): return jsonify({'hata':'Bu işlem yalnızca yönetici Dilek Kaya tarafından yapılabilir.'}),403
    a = AltProje.query.get_or_404(aid)
    if Fatura.query.filter_by(alt_proje_id=aid).first() or FaturaProje.query.filter_by(alt_proje_id=aid).first():
        return jsonify({'hata':'Bu alt proje faturalarda kullanıldığı için silinemez; pasife alabilirsiniz.'}),400
    db.session.delete(a); db.session.commit()
    return jsonify({'basarili':True})

@app.route('/api/yonetim/alt-proje-toggle/<int:aid>', methods=['POST'])
@login_gerekli
def api_alt_proje_toggle(aid):
    a = AltProje.query.get_or_404(aid); a.aktif = request.json.get('aktif', not a.aktif)
    db.session.commit(); return jsonify({'basarili':True})

# ─── EXCEL ŞABLON & IMPORT ──────────────────────────────────────────────────

@app.route('/api/kullanici-sablon-indir')
@login_gerekli
def kullanici_sablon_indir():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from flask import make_response; import io
    wb = Workbook()
    hf=Font(name='Arial',bold=True,color='FFFFFF',size=11)
    hfill=PatternFill(fill_type='solid',fgColor='2E3249')
    inf=Font(name='Arial',size=10,color='0000FF')
    efill=PatternFill(fill_type='solid',fgColor='EEF2FF')
    brd=Border(left=Side(style='thin',color='CCCCCC'),right=Side(style='thin',color='CCCCCC'),
               top=Side(style='thin',color='CCCCCC'),bottom=Side(style='thin',color='CCCCCC'))
    ws1=wb.active; ws1.title='Departmanlar'
    for col,bas in enumerate(['DEPARTMAN_KOD','DEPARTMAN_AD'],1):
        c=ws1.cell(1,col,bas); c.font=hf; c.fill=hfill; c.alignment=Alignment(horizontal='center'); c.border=brd
    for i,(kod,ad) in enumerate([('ASA','Aftermarket Satis'),('ESA','Endustriyel Satis'),('SSA','Servis'),('IDR','Idari Isler'),('MUH','Muhasebe')],2):
        for col,val in enumerate([kod,ad],1):
            c=ws1.cell(i,col,val); c.font=inf; c.fill=efill; c.border=brd
    ws1.column_dimensions['A'].width=20; ws1.column_dimensions['B'].width=35
    ws2=wb.create_sheet('Kullanicilar')
    for col,bas in enumerate(['AD_SOYAD','DEPARTMAN_KOD'],1):
        c=ws2.cell(1,col,bas); c.font=hf; c.fill=hfill; c.alignment=Alignment(horizontal='center'); c.border=brd
    for i,(ad,dept) in enumerate([('Ali Veli','ASA'),('Ayse Kaya','ESA')],2):
        for col,val in enumerate([ad,dept],1):
            c=ws2.cell(i,col,val); c.font=inf; c.fill=efill; c.border=brd
    ws2.column_dimensions['A'].width=30; ws2.column_dimensions['B'].width=20
    out=io.BytesIO(); wb.save(out); out.seek(0)
    resp=make_response(out.read())
    resp.headers['Content-Type']='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    resp.headers['Content-Disposition']='attachment; filename=kullanici_departman_sablon.xlsx'
    return resp

@app.route('/api/kullanici-excel-import', methods=['POST'])
@login_gerekli
def kullanici_excel_import():
    if 'dosya' not in request.files: return jsonify({'hata':'Dosya bulunamadı'}),400
    dosya=request.files['dosya']
    try:
        from openpyxl import load_workbook; import io
        wb=load_workbook(io.BytesIO(dosya.read()),data_only=True)
    except Exception as e: return jsonify({'hata':str(e)}),400
    dept_eklenen=kul_eklenen=atlanan=0
    if 'Departmanlar' in wb.sheetnames:
        for row in wb['Departmanlar'].iter_rows(min_row=2,values_only=True):
            kod=str(row[0]).strip().upper() if row[0] else ''; ad=str(row[1]).strip() if row[1] else ''
            if not kod or not ad or kod=='NONE': atlanan+=1; continue
            if not Departman.query.filter_by(kod=kod).first(): db.session.add(Departman(kod=kod,ad=ad)); dept_eklenen+=1
            else: atlanan+=1
    db.session.commit()
    if 'Kullanicilar' in wb.sheetnames:
        for row in wb['Kullanicilar'].iter_rows(min_row=2,values_only=True):
            ad=str(row[0]).strip() if row[0] else ''; dk=str(row[1]).strip().upper() if row[1] else ''
            if not ad or ad=='NONE': atlanan+=1; continue
            dept=Departman.query.filter_by(kod=dk).first() if dk else None
            db.session.add(Kullanici(ad_soyad=ad,departman_id=dept.id if dept else None)); kul_eklenen+=1
    db.session.commit()
    return jsonify({'basarili':True,'dept_eklenen':dept_eklenen,'kul_eklenen':kul_eklenen,'atlanan':atlanan})

@app.route('/api/proje-sablon-indir')
@login_gerekli
def proje_sablon_indir():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from flask import make_response; import io
    wb=Workbook()
    hf=Font(name='Arial',bold=True,color='FFFFFF',size=11)
    hfill=PatternFill(fill_type='solid',fgColor='2E3249')
    inf=Font(name='Arial',size=10,color='0000FF')
    efill=PatternFill(fill_type='solid',fgColor='EEF2FF')
    brd=Border(left=Side(style='thin',color='CCCCCC'),right=Side(style='thin',color='CCCCCC'),
               top=Side(style='thin',color='CCCCCC'),bottom=Side(style='thin',color='CCCCCC'))
    ws1=wb.active; ws1.title='Ana Projeler'
    for col,bas in enumerate(['ANA_PROJE_KOD','ANA_PROJE_AD'],1):
        c=ws1.cell(1,col,bas); c.font=hf; c.fill=hfill; c.alignment=Alignment(horizontal='center'); c.border=brd
    for i,(k,a) in enumerate([('GENEL','Genel Giderler'),('ASA-2026','Aftermarket 2026')],2):
        for col,val in enumerate([k,a],1): c=ws1.cell(i,col,val); c.font=inf; c.fill=efill; c.border=brd
    ws1.column_dimensions['A'].width=22; ws1.column_dimensions['B'].width=36
    ws2=wb.create_sheet('Alt Projeler')
    for col,bas in enumerate(['ANA_PROJE_KOD','ALT_PROJE_KOD','ALT_PROJE_AD'],1):
        c=ws2.cell(1,col,bas); c.font=hf; c.fill=hfill; c.alignment=Alignment(horizontal='center'); c.border=brd
    for i,(ana,alt,ad) in enumerate([('GENEL','KIR','Kirtasiye'),('ASA-2026','PAZ','Pazarlama')],2):
        for col,val in enumerate([ana,alt,ad],1): c=ws2.cell(i,col,val); c.font=inf; c.fill=efill; c.border=brd
    ws2.column_dimensions['A'].width=22; ws2.column_dimensions['B'].width=18; ws2.column_dimensions['C'].width=32
    out=io.BytesIO(); wb.save(out); out.seek(0)
    resp=make_response(out.read())
    resp.headers['Content-Type']='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    resp.headers['Content-Disposition']='attachment; filename=proje_kodu_sablon.xlsx'
    return resp

@app.route('/api/proje-excel-import', methods=['POST'])
@login_gerekli
def proje_excel_import():
    if 'dosya' not in request.files: return jsonify({'hata':'Dosya bulunamadı'}),400
    dosya=request.files['dosya']
    try:
        from openpyxl import load_workbook; import io
        wb=load_workbook(io.BytesIO(dosya.read()),data_only=True)
    except Exception as e: return jsonify({'hata':str(e)}),400
    ana_eklenen=alt_eklenen=atlanan=0
    if 'Ana Projeler' in wb.sheetnames:
        for row in wb['Ana Projeler'].iter_rows(min_row=2,values_only=True):
            kod=str(row[0]).strip().upper() if row[0] else ''; ad=str(row[1]).strip() if row[1] else ''
            if not kod or not ad or kod=='NONE': atlanan+=1; continue
            if not AnaProje.query.filter_by(kod=kod).first(): db.session.add(AnaProje(kod=kod,ad=ad)); ana_eklenen+=1
            else: atlanan+=1
    db.session.commit()
    if 'Alt Projeler' in wb.sheetnames:
        for row in wb['Alt Projeler'].iter_rows(min_row=2,values_only=True):
            ak=str(row[0]).strip().upper() if row[0] else ''; altk=str(row[1]).strip().upper() if row[1] else ''; ad=str(row[2]).strip() if row[2] else ''
            if not ak or not altk or not ad or ak=='NONE': atlanan+=1; continue
            ana=AnaProje.query.filter_by(kod=ak).first()
            # Bazı şablonlarda ana proje yalnızca Alt Projeler sayfasında yer
            # alır. Kaydı atlamak yerine ana projeyi aynı adla otomatik oluştur.
            if not ana:
                ana=AnaProje(kod=ak,ad=ak)
                db.session.add(ana); db.session.flush(); ana_eklenen+=1
            # Aynı kategori kodu altında farklı açıklamalar ayrı alt projedir.
            if not AltProje.query.filter_by(kod=altk,ad=ad,ana_proje_id=ana.id).first():
                db.session.add(AltProje(kod=altk,ad=ad,ana_proje_id=ana.id)); alt_eklenen+=1
            else: atlanan+=1
    standart_alt_kirilimlari_ekle(commit=False)
    db.session.commit()
    return jsonify({'basarili':True,'ana_eklenen':ana_eklenen,'alt_eklenen':alt_eklenen,'atlanan':atlanan})

@app.route('/static/sw.js')
def service_worker():
    from flask import send_from_directory, make_response
    resp = make_response(send_from_directory('static', 'sw.js'))
    resp.headers['Content-Type'] = 'application/javascript'
    resp.headers['Service-Worker-Allowed'] = '/'
    return resp

@app.route('/hakkinda')
@login_gerekli
def hakkinda():
    return render_template('hakkinda.html', oturum_kullanici=aktif_kullanici())

@app.route('/api/oturum-bilgisi')
def oturum_bilgisi():
    k = aktif_kullanici()
    if not k: return jsonify({}), 401
    return jsonify({'id': k.id, 'ad_soyad': k.ad_soyad, 'kisaltma': k.kisaltma,
                    'departman_kod': k.departman.kod if k.departman else '',
                    'fatura_silebilir': mim_fatura_silme_yetkisi()})

# Uygulama başlarken DB oluştur
with app.app_context():
    os.makedirs('database', exist_ok=True)
    os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
    try:
        db.create_all()
        try:
            db.session.execute(db.text(
                'ALTER TABLE kullanici ADD COLUMN email VARCHAR(254)'
            ))
            db.session.commit()
        except Exception:
            db.session.rollback()
        try:
            db.session.execute(db.text(
                'ALTER TABLE kullanici ADD COLUMN atanan_faturalari_gizle BOOLEAN NOT NULL DEFAULT FALSE'
            ))
            db.session.commit()
        except Exception:
            db.session.rollback()
        # Gerçek proje kodları ve açıklamaları eski 20/100 karakter
        # sınırından uzun olabildiği için PostgreSQL kolonlarını genişlet.
        try:
            for sql in (
                'ALTER TABLE ana_proje ALTER COLUMN kod TYPE VARCHAR(120)',
                'ALTER TABLE ana_proje ALTER COLUMN ad TYPE VARCHAR(300)',
                'ALTER TABLE alt_proje ALTER COLUMN kod TYPE VARCHAR(120)',
                'ALTER TABLE alt_proje ALTER COLUMN ad TYPE VARCHAR(300)',
            ):
                db.session.execute(db.text(sql))
            db.session.commit()
        except Exception:
            db.session.rollback()
        kolonlar = {c['name'] for c in inspect(db.engine).get_columns('fatura')}
        if 'atama_tamamlandi' not in kolonlar:
            db.session.execute(text(
                'ALTER TABLE fatura ADD COLUMN atama_tamamlandi BOOLEAN NOT NULL DEFAULT FALSE'
            ))
            db.session.commit()
        seed_data()
        standart_alt_kirilimlari_ekle()
    except Exception as e:
        print(f"Startup hatasi (devam ediyor): {e}")

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
